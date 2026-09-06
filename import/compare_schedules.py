#!/usr/bin/env python3
"""
Сравнение пофамильного расписания (docx) с основным расписанием (xlsx).
Извлекает данные из обоих файлов и находит расхождения.
"""
import re
import sys
from pathlib import Path
from collections import defaultdict

try:
    from docx import Document
except ImportError:
    print("Установка python-docx...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    from docx import Document

try:
    from openpyxl import load_workbook
except ImportError:
    print("Установка openpyxl...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    from openpyxl import load_workbook


def parse_weeks(s: str) -> list[int]:
    """Парсит строки недель: '1-7,12-17' -> [1,2,3,4,5,6,7,12,13,14,15,16,17]"""
    weeks = set()
    for part in s.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            pieces = part.split("-", 1)
            try:
                weeks.update(range(int(pieces[0]), int(pieces[1]) + 1))
            except ValueError:
                pass
        else:
            try:
                weeks.add(int(part))
            except ValueError:
                pass
    return sorted(weeks)


def extract_teacher_schedule_from_docx(filepath: str) -> dict:
    """
    Извлекает расписание преподавателей из docx файла.
    Возвращает: {teacher_name: {pair_num: {day_idx: [(group, weeks), ...]}}}
    """
    doc = Document(filepath)
    teacher_schedules = {}
    current_teacher = None

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if all(not c for c in cells):
                continue

            first_cell = cells[0] if cells else ""

            # Имя преподавателя — заглавные буквы, длинное, не цифра
            if (first_cell
                and first_cell.isupper()
                and len(first_cell) > 5
                and not first_cell.isdigit()
                and any(c.isalpha() for c in first_cell)):
                current_teacher = first_cell
                if current_teacher not in teacher_schedules:
                    teacher_schedules[current_teacher] = {}
                continue

            if not current_teacher:
                continue

            pair_text = first_cell
            if not pair_text or not pair_text.isdigit():
                continue
            pair_num = int(pair_text)
            if pair_num < 1 or pair_num > 8:
                continue

            if current_teacher not in teacher_schedules:
                teacher_schedules[current_teacher] = {}
            if pair_num not in teacher_schedules[current_teacher]:
                teacher_schedules[current_teacher][pair_num] = {}

            for day_idx in range(5):
                cell_idx = day_idx + 1
                if cell_idx >= len(cells):
                    break
                cell_text = cells[cell_idx]
                if not cell_text:
                    continue

                # Парсим: ГРУППА(недели) ГРУППА2(недели2) ...
                entries = []
                for m in re.finditer(r'([А-Яа-яёЁ\-]+(?:\d+)(?:-[А-Яа-яёЁ\-]*\d+)?)\(([^)]+)\)', cell_text):
                    group = m.group(1)
                    weeks = parse_weeks(m.group(2))
                    entries.append((group, weeks))

                if entries:
                    teacher_schedules[current_teacher][pair_num][day_idx] = entries

    return teacher_schedules


def extract_main_schedule_from_xlsx(filepath: str) -> dict:
    """
    Извлекает расписание из основного xlsx файла.
    Возвращает: {teacher_name: {(group, day_idx, pair_num): [weeks]}}
    """
    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active

    groups = []
    for col in range(3, ws.max_column + 1):
        val = ws.cell(5, col).value
        if val and str(val).strip():
            groups.append((col, str(val).strip()))

    day_map = {
        "ПОНЕДЕЛЬНИК": 0, "ВТОРНИК": 1, "СРЕДА": 2,
        "ЧЕТВЕРГ": 3, "ПЯТНИЦА": 4, "СУББОТА": 5,
    }

    # teacher -> {(group, day, pair): [weeks]}
    schedule_by_teacher = defaultdict(lambda: defaultdict(list))

    for row in range(6, ws.max_row + 1):
        day_cell = ws.cell(row, 1).value
        pair_cell = ws.cell(row, 2).value
        if not day_cell or not pair_cell:
            continue

        day_str = str(day_cell).strip().upper()
        day_idx = day_map.get(day_str)
        if day_idx is None:
            continue

        try:
            pair_num = int(pair_cell)
        except (ValueError, TypeError):
            continue

        for col, group_name in groups:
            cell_val = ws.cell(row, col).value
            if not cell_val:
                continue
            cell_text = str(cell_val).strip()
            if not cell_text:
                continue

            # Каждая ячейка может содержать несколько записей
            # Формат: аудитория предмет(недели) преподаватель
            # Разбиваем по паттерну начала записи
            parts = re.split(r'(?=(?:с\.з\.|ч\.з\.|[А-Яа-яёЁ]*\d))', cell_text)

            for part in parts:
                part = part.strip()
                if not part:
                    continue

                # Извлекаем недели
                week_matches = re.findall(r'\(([^)]+)\)', part)
                weeks = []
                for wm in week_matches:
                    weeks.extend(parse_weeks(wm))

                if not weeks:
                    continue

                # Извлекаем преподавателя — последняя часть строки кириллицей
                teacher_match = re.search(
                    r'([А-Яа-яёЁ][а-яёЁ]+\s+[А-Яа-яёЁ]\.?\s*[А-Яа-яёЁ]?\.)\s*$',
                    part
                )
                if teacher_match:
                    teacher_name = teacher_match.group(1).upper().strip()
                    # Нормализуем имя преподавателя
                    teacher_name = re.sub(r'\s+', ' ', teacher_name)
                    schedule_by_teacher[teacher_name][(group_name, day_idx, pair_num)].extend(weeks)

    wb.close()
    return schedule_by_teacher


def normalize_teacher_name(name: str) -> str:
    """Нормализует имя преподавателя для сравнения."""
    name = name.upper().strip()
    name = re.sub(r'\s+', ' ', name)
    # Убираем точки в инициалах для сравнения
    name_no_dots = name.replace('.', '').replace(' ', '')
    return name_no_dots


def compare_schedules(teacher_docx: dict, main_by_teacher: dict) -> list[dict]:
    """Сравнивает расписания и возвращает список расхождений."""
    differences = []

    # Создаём маппинг нормализованных имен
    main_name_map = {}
    for name in main_by_teacher:
        main_name_map[normalize_teacher_name(name)] = name

    for teacher_name, docx_schedule in teacher_docx.items():
        norm = normalize_teacher_name(teacher_name)
        matched_teacher = main_name_map.get(norm)

        if not matched_teacher:
            # Попробуем частичное совпадение
            for main_norm, main_name in main_name_map.items():
                if norm in main_norm or main_norm in norm:
                    matched_teacher = main_name
                    break

        if not matched_teacher:
            differences.append({
                "type": "TEACHER_NOT_FOUND",
                "teacher": teacher_name,
                "message": f"Преподаватель '{teacher_name}' не найден в основном расписании",
            })
            continue

        # Собираем данные из docx: {(group, day): {pair: [weeks]}}
        docx_data = defaultdict(lambda: defaultdict(lambda: defaultdict(set)))
        for pair_num, days in docx_schedule.items():
            for day_idx, entries in days.items():
                for group, weeks in entries:
                    for w in weeks:
                        docx_data[group][(day_idx, pair_num)].add(w)

        # Собираем данные из основного расписания
        main_entries = main_by_teacher[matched_teacher]
        main_data = defaultdict(lambda: defaultdict(lambda: defaultdict(set)))
        for (group, day_idx, pair_num), weeks in main_entries.items():
            for w in weeks:
                main_data[group][(day_idx, pair_num)].add(w)

        # Сравниваем группы
        docx_groups = set(docx_data.keys())
        main_groups = set(main_data.keys())

        for g in docx_groups - main_groups:
            details = []
            for (d, p), w in docx_data[g].items():
                day_name = ["Пн", "Вт", "Ср", "Чт", "Пт"][d]
                details.append(f"{day_name} п.{p} нед.{sorted(w)}")
            differences.append({
                "type": "GROUP_ONLY_DOCX",
                "teacher": teacher_name,
                "group": g,
                "message": f"  {teacher_name}: группа {g} ТОЛЬКО в docx ({'; '.join(details[:3])})",
            })

        for g in main_groups - docx_groups:
            details = []
            for (d, p), w in main_data[g].items():
                day_name = ["Пн", "Вт", "Ср", "Чт", "Пт"][d]
                details.append(f"{day_name} п.{p} нед.{sorted(w)}")
            differences.append({
                "type": "GROUP_ONLY_MAIN",
                "teacher": teacher_name,
                "group": g,
                "message": f"  {teacher_name}: группа {g} ТОЛЬКО в основном ({'; '.join(details[:3])})",
            })

        # Сравниваем недели для общих групп
        for g in docx_groups & main_groups:
            all_keys = set(docx_data[g].keys()) | set(main_data[g].keys())
            for key in all_keys:
                docx_weeks = docx_data[g].get(key, set())
                main_weeks = main_data[g].get(key, set())

                only_docx = docx_weeks - main_weeks
                only_main = main_weeks - docx_weeks

                if only_docx or only_main:
                    d, p = key
                    day_name = ["Пн", "Вт", "Ср", "Чт", "Пт"][d]
                    parts = []
                    if only_docx:
                        parts.append(f"docx: {sorted(only_docx)}")
                    if only_main:
                        parts.append(f"основное: {sorted(only_main)}")
                    differences.append({
                        "type": "WEEKS_DIFF",
                        "teacher": teacher_name,
                        "group": g,
                        "message": f"  {teacher_name}, {g}, {day_name} п.{p}: {'; '.join(parts)}",
                    })

    return differences


def main():
    base_dir = Path(__file__).parent
    docx_file = base_dir / "schedule" / "Пофамильное расписание 1 полугодие.docx"
    xlsx_file = base_dir / "schedule" / "Расписание.xlsx"

    if not docx_file.exists():
        print(f"ОШИБКА: {docx_file}")
        sys.exit(1)
    if not xlsx_file.exists():
        print(f"ОШИБКА: {xlsx_file}")
        sys.exit(1)

    print(f"DOCX: {docx_file.name}")
    teacher_docx = extract_teacher_schedule_from_docx(str(docx_file))
    print(f"  Преподавателей: {len(teacher_docx)}")

    print(f"XLSX: {xlsx_file.name}")
    main_by_teacher = extract_main_schedule_from_xlsx(str(xlsx_file))
    print(f"  Преподавателей: {len(main_by_teacher)}")

    print("\nСравнение...")
    diffs = compare_schedules(teacher_docx, main_by_teacher)

    if not diffs:
        print("\nРасписания совпадают!")
        return

    print(f"\nРасхождений: {len(diffs)}")

    by_type = defaultdict(list)
    for d in diffs:
        by_type[d["type"]].append(d)

    if "TEACHER_NOT_FOUND" in by_type:
        print(f"\nНе найдены в основном ({len(by_type['TEACHER_NOT_FOUND'])}):")
        for d in by_type["TEACHER_NOT_FOUND"]:
            print(d["message"])

    if "GROUP_ONLY_DOCX" in by_type:
        print(f"\nГруппы ТОЛЬКО в docx ({len(by_type['GROUP_ONLY_DOCX'])}):")
        for d in by_type["GROUP_ONLY_DOCX"]:
            print(d["message"])

    if "GROUP_ONLY_MAIN" in by_type:
        print(f"\nГруппы ТОЛЬКО в основном ({len(by_type['GROUP_ONLY_MAIN'])}):")
        for d in by_type["GROUP_ONLY_MAIN"]:
            print(d["message"])

    if "WEEKS_DIFF" in by_type:
        print(f"\nРазличия недель ({len(by_type['WEEKS_DIFF'])}):")
        for d in by_type["WEEKS_DIFF"]:
            print(d["message"])


if __name__ == "__main__":
    main()
